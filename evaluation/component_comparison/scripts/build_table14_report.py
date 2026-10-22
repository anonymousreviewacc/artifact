from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import csv, pathlib, json
BASE=pathlib.Path(r"/path/to/Research_project\Baseline_app_comparision")
CALC=BASE/"04_Calculations"
OUT=BASE/"Table14_Component_Results_and_Evidence.docx"
def readcsv(name):
    with (CALC/name).open(encoding="utf-8-sig",newline="") as f:return list(csv.DictReader(f))
summary=readcsv("Table14_Recalculated_Results.csv")
class_rows=readcsv("Table14_ECR_By_Class.csv")
meta=json.loads((BASE/"06_Documentation"/"Table14_Experiment_Metadata.json").read_text(encoding="utf-8"))

NAVY="17365D"; BLUE="1F4E79"; PALE="D9EAF7"; LIGHT="F1F5F9"; GRAY="596579"; WHITE="FFFFFF"
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd");shd.set(qn("w:fill"),fill);tcPr.append(shd)
def set_cell_margins(cell,top=95,start=90,bottom=95,end=90):
    tc=cell._tc;tcPr=tc.get_or_add_tcPr();mar=tcPr.first_child_found_in("w:tcMar")
    if mar is None: mar=OxmlElement("w:tcMar");tcPr.append(mar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=mar.find(qn("w:"+m))
        if node is None:node=OxmlElement("w:"+m);mar.append(node)
        node.set(qn("w:w"),str(v));node.set(qn("w:type"),"dxa")
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();el=OxmlElement("w:tblHeader");el.set(qn("w:val"),"true");trPr.append(el)
def set_repeat_and_keep(table):
    repeat_header(table.rows[0])
    for row in table.rows:
        for cell in row.cells: cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;set_cell_margins(cell)
def set_widths(table,widths):
    for row in table.rows:
        for i,w in enumerate(widths):row.cells[i].width=Inches(w)
def fmtpct(s):
    if s in ("", "N/A"):return s
    return f"{float(s)*100:.2f}%"
def fmt4(s):
    if s in ("", "N/A"):return s
    return f"{float(s):.4f}"
def style_table(table,header_fill=NAVY,font_size=9.2):
    table.style="Table Grid";table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False;set_repeat_and_keep(table)
    for i,row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i and cell==row.cells[0] else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name="Arial";run.font.size=Pt(font_size)
                    if i==0:run.font.bold=True;run.font.color.rgb=RGBColor.from_string(WHITE)
            if i==0:shade(cell,header_fill)
            elif i%2==0:shade(cell,LIGHT)

def add_bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet");p.paragraph_format.space_after=Pt(3);p.add_run(text)

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.64);sec.right_margin=Inches(.64)
normal=doc.styles["Normal"];normal.font.name="Arial";normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor.from_string("202833");normal.paragraph_format.space_after=Pt(6)
for sn in ("Heading 1","Heading 2"):
    doc.styles[sn].font.name="Arial";doc.styles[sn].font.color.rgb=RGBColor.from_string(NAVY)
    doc.styles[sn].paragraph_format.keep_with_next=True
doc.styles["Heading 1"].font.size=Pt(16);doc.styles["Heading 1"].font.bold=True
doc.styles["Heading 2"].font.size=Pt(12);doc.styles["Heading 2"].font.bold=True
# Title and summary
title_style=doc.styles["Title"]
title_style.font.color.rgb=RGBColor(0,0,0)
title_style.font.underline=False
p=doc.add_paragraph(style="Title")
p.paragraph_format.space_after=Pt(3)
pPr=p._p.get_or_add_pPr()
border=pPr.find(qn("w:pBdr"))
if border is not None:pPr.remove(border)
stylePr=title_style.element.find(qn("w:pPr"))
if stylePr is not None:
    styleBorder=stylePr.find(qn("w:pBdr"))
    if styleBorder is not None:stylePr.remove(styleBorder)
r=p.add_run("Table 14 Component Rerun Results")
r.font.name="Arial"
r.font.size=Pt(24)
r.font.bold=True
r.font.color.rgb=RGBColor(0,0,0)
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(12);rr=p.add_run("Matched generation ablations and local benchmark evidence  |  19 October 2026");rr.font.size=Pt(10);rr.font.color.rgb=RGBColor.from_string(GRAY)
full=next(x for x in summary if x["Configuration"]=="full_hybrid_density")
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(9)
p.add_run("The four generation configurations have now been rerun from the current generator source. ")
p.add_run(f"The full hybrid retained {full['Retained_candidate_rows']} candidates from {full['Attempts']} attempts, giving PSR {fmtpct(full['PSR'])}; its IMR was {fmtpct(full['IMR'])} ({full['Invalid_mutation_attempts']}/{full['Mutation_attempts']} mutation attempts), and its ECR was {fmtpct(full['ECR'])} ({full['Confirmed_trials']}/{full['Eligible_ECR_trials']} eligible trials). ")
p.add_run("The classifier-only macro-F1 remains 0.8645 on the existing 14,203-row family-grouped holdout. The generated-output F1 values in the other rows measure agreement between each requested generation class and the model prediction, so they use a different evaluation set.")

# Main table
h=doc.add_heading("Table 14 rerun values",level=1)
t=doc.add_table(rows=1,cols=6)
headers=["Configuration","PSR","IMR","ECR","Category macro-F1","Token-bigram diversity"]
for i,x in enumerate(headers):t.rows[0].cells[i].text=x
for row in summary:
    name={"rule_only":"Rule only","mutation_only_fixed2":"Mutation only","rule_mutation_density_no_filter":"Rule + mutation + density guidance, ungated","classifier_only":"Classifier only","full_hybrid_density":"Full hybrid density guided"}[row["Configuration"]]
    if row["Configuration"]=="classifier_only":
        cells=[name,"Not applicable","Not applicable","Not applicable",f"{fmt4(row['Generated_output_macro_F1'])}\n14,203 held-out rows","Not applicable"]
    else:
        imr="Not applicable\nno mutation path" if row["IMR"]=="" else f"{fmtpct(row['IMR'])}\n{row['Invalid_mutation_attempts']}/{row['Mutation_attempts']}"
        if row["Configuration"]=="rule_mutation_density_no_filter":imr += "\npost hoc"
        cells=[name,f"{fmtpct(row['PSR'])}\n{row['Valid_unique_per_call']}/{row['Attempts']}",imr,f"{fmtpct(row['ECR'])}\n{row['Confirmed_trials']}/{row['Eligible_ECR_trials']}",f"{fmt4(row['Generated_output_macro_F1'])}\n{row['Prediction_rows']} outputs",f"{float(row['Token_bigram_diversity']):.4f}\n{row['Retained_candidate_rows']} output rows"]
    cells_out=t.add_row().cells
    for i,value in enumerate(cells):cells_out[i].text=value
style_table(t,font_size=8.6);set_widths(t,[1.63,1.02,1.02,1.02,1.15,1.32])
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(8)
r=p.add_run("Reading the table. ");r.bold=True
p.add_run("PSR is the count of distinct structurally valid outputs retained per category and profile call divided by all generation attempts. IMR is structurally invalid mutation attempts divided by mutation attempts. ECR is behaviorally confirmed trials divided by eligible trials on the listed local sinks. Diversity follows the manuscript’s adjacent-token-bigram Jaccard definition. The ungated arm emits every attempt; its invalid-mutation ratio is calculated after generation and its PSR numerator is also calculated post hoc.")

# Design and procedure
h=doc.add_heading("Configurations and generation procedure",level=1)
p=doc.add_paragraph("Each generation arm ran 60 attempts for each of 12 vulnerability classes and three profiles, for 2,160 attempts per configuration. The Android emulator executed an isolated debug build copied into the baseline experiment folder. Its copied GNNGenerator.java has the same SHA-256 hash as the production source. The test build used one seeded Random stream per configuration, class, and profile, matching the production call order for seed selection, path choice, and mutation-operator draws. The production app and manuscript source were not changed.")
configs=[
("Rule only","Select a class-specific seed directly, apply the selected profile transform, then apply the structural predicate and per-call duplicate suppression. No mutation path is used."),
("Mutation only","Force the mutation path for every attempt at fixed moderate intensity, which applies two class-compatible mutation operations, then apply the profile transform, structural predicate, and per-call duplicate suppression."),
("Rule + mutation + density guidance, ungated","Use the production 70/30 direct and mutation choice and density-based mutation intensity, but emit every attempt without structural filtering or duplicate suppression. This isolates the output gate; PSR and IMR are evaluated post hoc."),
("Full hybrid density guided","Use the production 70/30 path choice and density-based intensity, followed by the app’s structural predicate and per-call duplicate suppression."),
("Classifier only","Recalculate category macro-F1 from the existing family-grouped held-out predictions. This row has no generation or runtime trial metrics; the classifier was not retrained.")]
for label,body in configs:
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5);r=p.add_run(label+". ");r.bold=True;p.add_run(body)

# Metric definitions
h=doc.add_heading("Metric definitions",level=1)
defs=[
("PSR","Distinct structurally valid payload entries retained within each category/profile call, summed across calls, divided by 2,160 attempts. For the ungated arm, the same predicate is applied after emission."),
("IMR","Mutation attempts failing the class-specific structural predicate divided by mutation attempts. Rule only is not applicable because it makes no mutation attempts. The ungated arm’s value is post hoc and does not mean the app rejected those candidates."),
("ECR","Confirmed trials divided by eligible trials. A trial is one emitted payload sent to one supported local target route with an available behavioral oracle. Repeated payloads in different attempts remain separate trials."),
("Generated-output category macro-F1","TFLite-predicted category compared with the category requested from the generator, scored on all emitted candidate rows for that configuration and macro-averaged across the 12 labels. This measures output characterization on generated candidates, not the classifier holdout."),
("Classifier-only category macro-F1","Macro-F1 recomputed from 14,203 rows in Model 6’s existing family-grouped held-out prediction file. The recalculated value is 0.8644696, rounded to 0.8645."),
("Token-bigram diversity","One minus the mean pairwise Jaccard similarity of adjacent-token bigram sets, using the app’s tokenizer v1. The calculation pools emitted output rows; duplicates emitted by the ungated arm remain in the set and therefore affect the pairwise score.")]
t=doc.add_table(rows=1,cols=2);t.rows[0].cells[0].text="Metric";t.rows[0].cells[1].text="Calculation used"
for a,b in defs:
    cells=t.add_row().cells;cells[0].text=a;cells[1].text=b
style_table(t,font_size=9.2);set_widths(t,[2.0,5.1])

# ECR by class
h=doc.add_heading("Behavioral confirmation scope",level=1)
p=doc.add_paragraph("All browser and API requests were restricted to the four local training targets. The ECR denominator includes six classes across 15 class/target pairs. The class-level table below pools the supported target routes for each class. Per-target counts and all unscored probe outcomes are in the accompanying CSV files.")
class_names={"rule_only":"Rule only","mutation_only_fixed2":"Mutation only","rule_mutation_density_no_filter":"Guided ungated","full_hybrid_density":"Full hybrid"}
class_order=["SQL Injection","XSS","Command Injection","Path Traversal","Local File Inclusion","Open Redirect"]
variants=["rule_only","mutation_only_fixed2","rule_mutation_density_no_filter","full_hybrid_density"]
lookup={(r["variant"],r["category"]):r for r in class_rows}
t=doc.add_table(rows=1,cols=5)
for i,v in enumerate(["Class","Rule only","Mutation only","Guided ungated","Full hybrid"]):t.rows[0].cells[i].text=v
for cat in class_order:
    cells=t.add_row().cells;cells[0].text=cat
    for i,v in enumerate(variants,1):
        x=lookup.get((v,cat));cells[i].text=(f"{int(x['confirmed_trials'])}/{int(x['eligible_trials'])}  ({float(x['ecr'])*100:.1f}%)" if x else "Not estimable")
style_table(t,font_size=8.8);set_widths(t,[1.82,1.30,1.30,1.35,1.33])

p=doc.add_paragraph();r=p.add_run("Excluded from ECR. ");r.bold=True
p.add_run("bWAPP SQL Injection was excluded because its configured route produced no usable SQL oracle. SSRF was not submitted because XVWA’s available handler fetches the URL and writes a random image file. RFI payloads were not sent to remote include sinks. IDOR, XXE, GraphQL Injection, and API Abuse were recorded as probes where routes existed, but their responses did not provide a validated confirmation oracle. These are not counted as negative ECR outcomes.")
p=doc.add_paragraph();r=p.add_run("Setup limitation. ");r.bold=True
p.add_run("The benchmark runner attempted local logins for DVWA, bWAPP, and XVWA. XVWA’s setup step timed out, so XVWA results should be described as unauthenticated unless the session is verified separately. Juice Shop was tested anonymously. Benchmark result files preserve the HTTP status, outcome, evidence code, elapsed time, and errors for each executed trial.")

# Observations and caveats
h=doc.add_heading("What the rerun shows",level=1)
for text in [
"Rule only retained the fewest candidates, 193 of 2,160 attempts. Mutation only retained 314. The full hybrid retained 338. The ungated density-guided arm emitted all 2,160 attempts, although only 338 distinct structurally valid outputs were found when the predicate and duplicate check were applied after generation.",
"The full hybrid ECR was 259/606 (42.74%). The ungated guided arm’s ECR was 1,072/2,700 (39.70%). These rates use different numbers of emitted candidates and count repeated payload attempts separately. They are descriptive one-run rates, not a causal estimate that filtering alone raised ECR.",
"The old manuscript values for rule-only, mutation-only, and the intermediate configuration were not recoverable from raw historical logs. They are retained only as historical values in the manuscript. Replace them with the rerun values only after deciding whether the one-run, fixed-seed comparison is sufficient for the paper’s claims.",
"The classifier-only held-out macro-F1 of 0.8645 is not the same measure as the generated-output macro-F1 values (0.4521 to 0.4807). The former uses labeled held-out corpus examples; the latter evaluates whether generated candidates are classified as the requested category.",
"Each configuration was run once. The experiment does not provide confidence intervals or repeated-run significance tests. The mutation-only arm uses a fixed moderate intensity as a defined ablation baseline; it does not claim to be an earlier app mode."
]:add_bullet(doc,text)

h=doc.add_heading("Files and evidence",level=1)
files=[
"Generation attempts, emitted payloads, predictions, and run manifest: 02_Generation_Results/android_run/final_rng_matched/baseline_ablation_20261019T024137.729884Z/",
"Per-configuration benchmark outcomes and metadata: 03_Benchmark_Results/shared_rng_final/",
"Benchmark input payload lists: 03_Benchmark_Results/inputs_shared_rng_final/",
"Summary calculations: 04_Calculations/Table14_Recalculated_Results.csv, Table14_ECR_By_Class.csv, Table14_ECR_By_Class_and_Target.csv, Table14_Generation_By_Class_Profile.csv, Table14_Generated_Prediction_Metrics.csv, and Table14_All_Benchmark_Outcomes.csv",
"Local benchmark screenshots: 05_Screenshots/shared_rng_final/",
"Run metadata and source hashes: 06_Documentation/Table14_Experiment_Metadata.json",
"The source harness and scripts are in 01_Test_Harness_Source/. The original project source and main_doc.tex were left unchanged."
]
for x in files:add_bullet(doc,x)
# Footer with page number
for section in doc.sections:
    fp=section.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=fp.add_run("Table 14 rerun evidence  |  ");run.font.name="Arial";run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string(GRAY)
    fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");fp._p.append(fld)
doc.core_properties.title="Table 14 Component Rerun Results"
doc.core_properties.subject="Matched generator ablations and local behavioral benchmark results"
doc.core_properties.author="PLGenN project"
doc.save(OUT)
print(OUT)

